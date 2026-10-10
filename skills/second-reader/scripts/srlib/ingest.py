"""Tabular ingest (CSV / TSV / XLSX / JSON) -> table_rows + column profile + role mapping suggestion."""
import csv
import io
import json
import re
import unicodedata
from collections import Counter
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

from .numparse import parse_date, parse_number

INGEST_VERSION = "2"
TABULAR = {".csv": "csv", ".tsv": "csv", ".xlsx": "xlsx", ".xlsm": "xlsx", ".json": "json"}
DISTINCT_CAP = 50000
SAMPLE_N = 5
SAMPLE_CLIP = 80
HEADER_SCAN_ROWS = 20


def detect_type(path):
    return TABULAR.get(Path(path).suffix.lower())


# ---- readers: yield (sheet, row_no, cells) with row_no = 1-based source row -----------------
def read_csv(path):
    raw = Path(path).read_bytes()
    for enc in ("utf-8-sig", "cp1252"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    else:
        text = raw.decode("latin-1")
    sample = text[:20000]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    for i, row in enumerate(csv.reader(io.StringIO(text, newline=""), dialect), start=1):
        yield "csv", i, row


def read_xlsx(path):
    try:
        import openpyxl
    except ImportError as exc:
        raise SystemExit("XLSX needs openpyxl: run via scripts/sr (uses `uv run --with openpyxl`) "
                         "or install openpyxl") from exc
    wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    try:
        for ws in wb.worksheets:
            for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
                yield ws.title, i, list(row)
    finally:
        wb.close()


def _flatten(obj, prefix=""):
    out = {}
    for k, v in obj.items():
        key = f"{prefix}{k}"
        if isinstance(v, dict):
            out.update(_flatten(v, key + "."))
        elif isinstance(v, list):
            out[key] = json.dumps(v, sort_keys=True) if v else None
        else:
            out[key] = v
    return out


def read_json(path):
    data = json.loads(Path(path).read_text(encoding="utf-8-sig"))
    if isinstance(data, dict):
        lists = [(k, v) for k, v in data.items() if isinstance(v, list) and v and isinstance(v[0], dict)]
        if not lists:
            data = [data]
        else:
            data = max(lists, key=lambda kv: len(kv[1]))[1]
    if not isinstance(data, list) or not all(isinstance(r, dict) for r in data):
        raise SystemExit("JSON must be an array of objects (or an object holding one)")
    flat = [_flatten(r) for r in data]
    cols = list(dict.fromkeys(k for r in flat for k in r))
    yield "json", 1, cols
    for i, r in enumerate(flat, start=2):
        yield "json", i, [r.get(c) for c in cols]


READERS = {"csv": read_csv, "xlsx": read_xlsx, "json": read_json}


def _jsonable(v):
    if isinstance(v, (datetime, date)):
        return v.isoformat()
    if isinstance(v, Decimal):
        return str(v)
    return v


def _is_header(cells, width):
    filled = sum(1 for c in cells if not _blank(c))
    return filled >= max(2, 0.6 * width)


def _pick_header(buf):
    """Index of the header row within the first rows of a sheet. Title/metadata rows above a table have few
    filled cells; the header is the first row with at least 60% of the widest row's cells filled."""
    width = max(sum(1 for c in cells if not _blank(c)) for _, cells in buf)
    for k, (_, cells) in enumerate(buf):
        if _is_header(cells, width):
            return k
    return 0


def _header_names(cells):
    names, seen = [], Counter()
    for j, c in enumerate(cells):
        name = str(c).strip() if c is not None and str(c).strip() else f"col{j + 1}"
        seen[name] += 1
        names.append(name if seen[name] == 1 else f"{name}_{seen[name]}")
    return names


def ingest_rows(db, doc_id, path, doc_type):
    """Loads rows into table_rows. Rows above the detected header are kept as a preamble (not data).
    Returns {sheet: {header_row, columns, data_rows, preamble}}."""
    db.execute("DELETE FROM table_rows WHERE doc_id=?", (doc_id,))
    headers, batch, data_rows = {}, [], Counter()
    pending = {}                       # sheet -> first non-empty rows, until the header is chosen

    def start_sheet(sheet, buf):
        k = _pick_header(buf)
        rn, cells = buf[k]
        headers[sheet] = {"header_row": rn, "columns": _header_names(cells),
                          "preamble": [[r, " | ".join(str(c) for c in cs if not _blank(c))[:SAMPLE_CLIP * 2]]
                                       for r, cs in buf[:k]]}
        for r, cs in buf[k + 1:]:
            add_row(sheet, r, cs)

    def add_row(sheet, row_no, cells):
        nonlocal batch
        batch.append((doc_id, sheet, row_no, json.dumps([_jsonable(c) for c in cells])))
        data_rows[sheet] += 1
        if len(batch) >= 5000:
            db.executemany("INSERT INTO table_rows VALUES(?,?,?,?)", batch)
            batch = []

    for sheet, row_no, cells in READERS[doc_type](path):
        if all(_blank(c) for c in cells):
            continue
        if sheet in headers:
            add_row(sheet, row_no, cells)
            continue
        buf = pending.setdefault(sheet, [])
        buf.append((row_no, cells))
        if len(buf) >= HEADER_SCAN_ROWS:
            start_sheet(sheet, pending.pop(sheet))
    for sheet, buf in pending.items():
        start_sheet(sheet, buf)
    if batch:
        db.executemany("INSERT INTO table_rows VALUES(?,?,?,?)", batch)
    for sheet in headers:
        headers[sheet]["data_rows"] = data_rows[sheet]
    return headers


# ---- profiling -------------------------------------------------------------------------------
def iter_rows(db, doc_id, sheet):
    for r in db.execute("SELECT row_no, data FROM table_rows WHERE doc_id=? AND sheet=? ORDER BY row_no",
                        (doc_id, sheet)):
        yield r["row_no"], json.loads(r["data"])


def _blank(v):
    return v is None or (isinstance(v, str) and not v.strip())


def profile_sheet(db, doc_id, sheet, columns):
    n = len(columns)
    stats = [{"n": 0, "nulls": 0, "numeric": 0, "date": 0, "text": 0, "ws_issues": 0, "distinct": set(),
              "overflow": False, "samples": [], "min": None, "max": None, "dmin": None, "dmax": None,
              "num_styles": Counter(), "dayfirst_votes": 0, "monthfirst_votes": 0} for _ in range(n)]
    rows = 0
    for _, cells in iter_rows(db, doc_id, sheet):
        rows += 1
        for j in range(n):
            v = cells[j] if j < len(cells) else None
            st = stats[j]
            st["n"] += 1
            if _blank(v):
                st["nulls"] += 1
                continue
            if isinstance(v, str) and v != v.strip():
                st["ws_issues"] += 1
            if len(st["distinct"]) < DISTINCT_CAP:
                st["distinct"].add(str(v).strip())
            else:
                st["overflow"] = True
            if len(st["samples"]) < SAMPLE_N and str(v)[:SAMPLE_CLIP] not in st["samples"]:
                st["samples"].append(str(v)[:SAMPLE_CLIP])
            num, nf = parse_number(v)
            d, df = parse_date(v) if not isinstance(v, (int, float)) else (None, set())
            if d is not None and (num is None or isinstance(v, str) and re.search(r"[-/:T]", v)):
                st["date"] += 1
                st["dmin"] = d if st["dmin"] is None or d < st["dmin"] else st["dmin"]
                st["dmax"] = d if st["dmax"] is None or d > st["dmax"] else st["dmax"]
                if isinstance(v, str):
                    m = re.match(r"^(\d{1,2})[-/.](\d{1,2})[-/.]", v.strip())
                    if m:
                        a, b = int(m.group(1)), int(m.group(2))
                        st["dayfirst_votes"] += a > 12
                        st["monthfirst_votes"] += b > 12
            elif num is not None:
                st["numeric"] += 1
                st["min"] = num if st["min"] is None or num < st["min"] else st["min"]
                st["max"] = num if st["max"] is None or num > st["max"] else st["max"]
                for f in nf & {"style_dot_decimal", "style_comma_decimal", "ambiguous_sep"}:
                    st["num_styles"][f] += 1
            else:
                st["text"] += 1
    cols = []
    for name, st in zip(columns, stats):
        filled = st["n"] - st["nulls"]
        if filled == 0:
            ctype = "empty"
        elif st["numeric"] / filled >= 0.95:
            ctype = "numeric"
        elif st["date"] / filled >= 0.95:
            ctype = "date"
        elif st["numeric"] + st["date"] and (st["numeric"] + st["date"]) / filled >= 0.5:
            ctype = "mixed"
        else:
            ctype = "text"
        cols.append({
            "name": name, "type": ctype, "filled": filled, "nulls": st["nulls"],
            "null_rate": round(st["nulls"] / st["n"], 4) if st["n"] else 0,
            "distinct": len(st["distinct"]), "distinct_capped": st["overflow"],
            "numeric": st["numeric"], "date": st["date"], "text": st["text"], "ws_issues": st["ws_issues"],
            "min": None if st["min"] is None else str(st["min"]), "max": None if st["max"] is None else str(st["max"]),
            "date_min": st["dmin"].isoformat() if st["dmin"] else None,
            "date_max": st["dmax"].isoformat() if st["dmax"] else None,
            "num_styles": dict(st["num_styles"]),
            "dayfirst": (True if st["dayfirst_votes"] > st["monthfirst_votes"]
                         else False if st["monthfirst_votes"] > st["dayfirst_votes"] else None),
            "samples": st["samples"]})
    return {"rows": rows, "columns": cols}


# ---- role mapping (heuristic; a human or Sonnet can override with `sr map --set role=Column`) ----
ROLE_SYNONYMS = {
    "qty": ["quantity", "qty", "menge", "anzahl", "stk", "quantite", "quantita", "cantidad", "units", "orderqty",
            "qtyordered", "pieces", "pcs"],
    "unit_price": ["unitprice", "price", "unitcost", "prixunitaire", "prix", "einzelpreis", "preis", "stuckpreis",
                   "preisstk", "prezzo", "prezzounitario", "precio", "preciounitario", "priceperunit", "rate"],
    "line_total": ["linetotal", "total", "lineamount", "extendedprice", "extended", "amount", "totalprice",
                   "linevalue", "netvalue", "gesamtpreis", "gesamt", "gesamtbetrag", "betrag", "summe", "montant",
                   "totale", "importe"],
    "date": ["date", "invoicedate", "orderdate", "quotedate", "datum", "bestelldatum", "belegdatum", "data",
             "fecha", "createdat", "created"],
    "date_due": ["deliverydate", "duedate", "needby", "needbydate", "requireddate", "shipdate", "promisedate",
                 "liefertermin", "lieferdatum", "wunschtermin", "datelivraison", "datadiconsegna"],
    "item_id": ["stockcode", "partno", "partnumber", "pn", "sku", "itemid", "itemcode", "article", "artikel",
                "artikelnr", "artikelnummer", "teilenummer", "materialnummer", "productcode", "codicearticolo",
                "item", "part", "mpn"],
    "description": ["description", "desc", "itemdescription", "name", "productname", "beschreibung", "bezeichnung",
                    "descrizione", "descripcion", "designation"],
    "doc_ref": ["invoiceno", "invoice", "orderno", "ordernumber", "po", "ponumber", "rfq", "quoteno", "reference",
                "bestellnummer", "belegnummer", "angebotsnummer"],
    "party": ["supplier", "vendor", "suppliername", "vendorname", "customer", "customerid", "buyer", "company",
              "lieferant", "kunde", "fournisseur", "fornitore", "proveedor"],
    "currency": ["currency", "curr", "ccy", "wahrung", "devise", "valuta", "moneda"],
    "lead_time": ["leadtime", "leaddays", "deliverytime", "lieferzeit", "delailivraison", "tempiconsegna"],
}


def _norm_header(h):
    h = unicodedata.normalize("NFKD", h.lower().replace("ß", "ss"))
    h = "".join(ch for ch in h if not unicodedata.combining(ch))
    return re.sub(r"[^a-z0-9]", "", h)


def suggest_mapping(profile_by_sheet):
    """-> (mapping, notes). Matching is by HEADER first (exact synonym = auto-exact, substring = auto-fuzzy).
    The column type only breaks ties and rejects columns that are clearly the wrong kind (text where a number is
    expected). A column with a matching header but dirty values (type `mixed`) is still mapped, flagged
    `dirty: true`, so the checks that detect the dirt still run. Header matches rejected for type are returned
    in `notes` so the brief can say why a role is missing."""
    want_type = {"qty": "numeric", "unit_price": "numeric", "line_total": "numeric", "date": "date",
                 "date_due": "date", "lead_time": "numeric"}
    mapping, notes = {}, []
    for sheet, prof in profile_by_sheet.items():
        taken = set()
        for strength in ("auto-exact", "auto-fuzzy"):
            for role, syns in ROLE_SYNONYMS.items():
                if role in mapping:
                    continue
                cands = []
                for col in prof["columns"]:
                    h = _norm_header(col["name"])
                    if col["name"] in taken or not h:
                        continue
                    if (h in syns) if strength == "auto-exact" else any(len(s) >= 4 and s in h for s in syns):
                        cands.append(col)
                if not cands:
                    continue
                wt = want_type.get(role)
                ok = [c for c in cands if not wt or c["type"] == wt] or \
                     [c for c in cands if wt and c["type"] == "mixed"]
                if not ok:
                    c = cands[0]
                    notes.append({"role": role, "col": c["name"], "sheet": sheet,
                                  "reason": f"header looks like {role} but the column is {c['type']}"})
                    continue
                col = ok[0]
                entry = {"sheet": sheet, "col": col["name"], "source": strength}
                if wt and col["type"] != wt:
                    entry["dirty"] = True
                mapping[role] = entry
                taken.add(col["name"])
    return mapping, notes
